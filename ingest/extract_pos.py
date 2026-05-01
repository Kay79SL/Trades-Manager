"""
extract_pos.py
==============
Phase 3 of DACARag: Purchase Order ingestion + LLM enrichment.

Two-stage processing for POs (different from emails because POs are
structured documents, not unstructured prose):

STAGE 1 — Parse:
  Reads the Excel PO workbook (one sheet per PO) using known cell positions
  and writes each PO as a document to MongoDB `pos` collection. This is
  deterministic structured parsing — no LLM needed.

STAGE 2 — LLM enrich:
  For each PO, calls Claude Haiku to add an `extracted` sub-document with
  interpretive fields that aren't directly in the cells:
    - summary                one-sentence description
    - complexity             low | medium | high
    - complexity_reasoning   why
    - risk_flags             list of concerns (or empty)
    - expected_duration_days rough estimate
    - materials_categories   broad categories (e.g. boiler-related, fittings, sanitaryware)
    - trade_inferred         plumber | carpenter | electrician
    - reasoning              one-sentence explanation

Saves to:
    pos collection            — one document per PO with all parsed fields
    pos.extracted (sub-doc)   — LLM analysis fields

Usage:
    python ingest\\extract_pos.py
    python ingest\\extract_pos.py --skip-parse   # skip Excel parse, only run LLM
    python ingest\\extract_pos.py --skip-llm     # skip LLM, only parse
    python ingest\\extract_pos.py --xlsx PATH    # path to PO workbook (default: data\\DACARag_Purchase_Order_Template.xlsx)
    python ingest\\extract_pos.py --force        # re-extract all
    python ingest\\extract_pos.py --dry-run      # print first prompt, no API calls
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path
from typing import Any

from anthropic import Anthropic
from dotenv import load_dotenv
from openpyxl import load_workbook
from pymongo import MongoClient
from tqdm import tqdm


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

DEFAULT_XLSX = Path("data") / "DACARag_Purchase_Order_Template.xlsx"
MODEL = "claude-haiku-4-5-20251001"
MAX_TOKENS = 800

# Sheets to skip (not POs)
SKIP_SHEETS = {"Empty Template", "How to use"}

SYSTEM_PROMPT_POS = """You are a purchase order analyst for a chatbot serving Irish trades.
You read parsed PO documents and add interpretive fields needed for downstream retrieval and triage.

You always respond with a single valid JSON object — no preamble, no commentary, no markdown
code fences. The JSON must conform exactly to the schema provided in the user message.

Be specific:
  - Avoid generic fillers like "good quality work" or "professional service".
  - Risk flags should only fire when there's a real concern (e.g. unusually small materials
    budget for the labour quoted, missing Eircode, suspicious total, very tight delivery date).
  - For complexity, weigh the number of items, total value, and trade discipline. A €144
    boiler service is "low". A €3,800 full house rewire is "high". A €2,000 wardrobe build is "medium".
  - For expected_duration_days, give a realistic on-site duration (1 day for service, 5-10 for
    rewires/kitchens, etc).
"""


# ---------------------------------------------------------------------------
# Stage 1 — Excel parsing
# ---------------------------------------------------------------------------

def parse_po_sheet(ws) -> dict[str, Any] | None:
    """Parse one PO sheet into a structured dict.

    Cell positions match the layout produced by build_po_template.py.
    Returns None if the sheet looks empty/template-only.
    """
    po_number = ws["E3"].value
    if not po_number or str(po_number).startswith("PO-YYYY"):
        return None  # empty template

    def _val(cell):
        v = ws[cell].value
        return v if v is not None else ""

    def _split_city_eircode(combined):
        """Split 'Dublin 12 · D12 X9F8' into ('Dublin 12', 'D12 X9F8')."""
        if not combined:
            return ("", "")
        s = str(combined)
        if "·" in s:
            parts = s.split("·", 1)
            return (parts[0].strip(), parts[1].strip())
        return (s.strip(), "")

    trade_city, trade_eircode = _split_city_eircode(_val("C9"))
    cust_city, cust_eircode = _split_city_eircode(_val("F9"))

    # Date is shown as "Date: YYYY-MM-DD" in E4
    date_raw = str(_val("E4")).replace("Date:", "").strip()

    po: dict[str, Any] = {
        "po_number":         po_number,
        "po_date":           date_raw,
        "sheet_name":        ws.title,

        # FROM (Trade)
        "trade_name":        _val("C7"),
        "trade_addr":        _val("C8"),
        "trade_city":        trade_city,
        "trade_eircode":     trade_eircode,
        "trade_email":       _val("C10"),
        "trade_phone":       _val("C11"),
        "trade_vat":         _val("C12"),

        # TO (Customer)
        "cust_name":         _val("F7"),
        "cust_addr":         _val("F8"),
        "cust_city":         cust_city,
        "cust_eircode":      cust_eircode,
        "cust_email":        _val("F10"),
        "cust_phone":        _val("F11"),
        "po_status":         _val("F12"),

        # Job overview
        "job_type":          _val("C15"),
        "job_desc":          _val("C16"),
        "delivery_date":     str(_val("F15")),
        "payment_terms":     _val("F16"),

        "line_items":        [],
    }

    # Read line items (rows 20+ until blank Item ID column)
    row = 20
    while True:
        item_id = ws.cell(row=row, column=2).value  # column B
        if not item_id or str(item_id).startswith("["):
            break
        po["line_items"].append({
            "item_id":           item_id,
            "description":       ws.cell(row=row, column=3).value,
            "quantity":          ws.cell(row=row, column=4).value,
            "unit":              ws.cell(row=row, column=5).value,
            "unit_price_ex_vat": ws.cell(row=row, column=6).value,
            "line_total_ex_vat": ws.cell(row=row, column=7).value,
        })
        row += 1

    last_item_row = row - 1
    if last_item_row < 20:
        return None  # no line items found

    # Totals start at last_item_row + 2 (one blank row gap)
    sub_row = last_item_row + 2
    po["materials_subtotal"] = ws.cell(row=sub_row,     column=7).value
    po["labour_cost"]        = ws.cell(row=sub_row + 1, column=7).value
    po["subtotal_ex_vat"]    = ws.cell(row=sub_row + 2, column=7).value
    po["vat_23pct"]          = ws.cell(row=sub_row + 3, column=7).value
    po["total_inc_vat"]      = ws.cell(row=sub_row + 4, column=7).value

    return po


def parse_workbook(xlsx_path: Path) -> list[dict]:
    """Parse all PO sheets from the workbook, skipping non-PO sheets."""
    if not xlsx_path.exists():
        print(f"ERROR: Excel file not found at {xlsx_path}")
        sys.exit(1)

    print(f"Loading workbook: {xlsx_path}")
    # data_only=True so we get calculated formula values (not the formula text)
    wb = load_workbook(str(xlsx_path), data_only=True)

    pos = []
    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS or sheet_name.startswith("How"):
            continue
        ws = wb[sheet_name]
        po = parse_po_sheet(ws)
        if po is not None:
            pos.append(po)
    return pos


def link_to_customer(po: dict, customers_by_email: dict) -> str | None:
    """Match the PO's cust_email back to a customer_id from the customers collection."""
    email = (po.get("cust_email") or "").strip().lower()
    if not email:
        return None
    matched = customers_by_email.get(email)
    if matched:
        return matched["customer_id"]
    return None


# ---------------------------------------------------------------------------
# Stage 2 — LLM enrichment
# ---------------------------------------------------------------------------

def build_po_prompt(po: dict) -> str:
    """Build the per-PO user prompt for LLM analysis."""
    lines_text = "\n".join(
        f"  - {li['item_id']:12} {li['description']:30} qty {li['quantity']} {li['unit']} "
        f"@ €{li['unit_price_ex_vat']:.2f} = €{li['line_total_ex_vat']:.2f}"
        for li in po["line_items"]
    )

    def fmt_money(v):
        try:
            return f"€{float(v):.2f}"
        except (ValueError, TypeError):
            return str(v)

    return f"""Analyse this purchase order:

PO_NUMBER:        {po['po_number']}
PO_DATE:          {po['po_date']}
TRADE:            {po['trade_name']}
CUSTOMER:         {po['cust_name']}
CUSTOMER_COUNTY:  {po['cust_city']}
JOB_TYPE:         {po['job_type']}
JOB_DESCRIPTION:  {po['job_desc']}
PAYMENT_TERMS:    {po['payment_terms']}
DELIVERY_DATE:    {po['delivery_date']}
STATUS:           {po['po_status']}

LINE ITEMS ({len(po['line_items'])} items):
{lines_text}

TOTALS:
  materials subtotal: {fmt_money(po.get('materials_subtotal'))}
  labour cost:        {fmt_money(po.get('labour_cost'))}
  subtotal ex VAT:    {fmt_money(po.get('subtotal_ex_vat'))}
  VAT 23%:            {fmt_money(po.get('vat_23pct'))}
  TOTAL inc VAT:      {fmt_money(po.get('total_inc_vat'))}

REQUIRED OUTPUT SCHEMA (respond with this JSON only — no fences, no commentary):
{{
  "summary":               "<one-sentence summary of what is being requested>",
  "complexity":            "low" | "medium" | "high",
  "complexity_reasoning":  "<why>",
  "risk_flags":            ["<flag 1>", "<flag 2>", ...],
  "expected_duration_days":<integer estimate of on-site days>,
  "materials_categories":  ["<category 1>", "<category 2>", ...],
  "trade_inferred":        "plumber" | "carpenter" | "electrician",
  "reasoning":             "<one-sentence explanation>"
}}
"""


def parse_json_response(text: str) -> dict[str, Any]:
    """Parse Claude's response, tolerating markdown fences."""
    text = text.strip()
    if text.startswith("```"):
        first_newline = text.find("\n")
        if first_newline != -1:
            text = text[first_newline + 1:]
        if text.endswith("```"):
            text = text[:-3]
        text = text.strip()
    return json.loads(text)


def extract_one(client: Anthropic, system_prompt: str, user_prompt: str) -> dict:
    response = client.messages.create(
        model=MODEL,
        max_tokens=MAX_TOKENS,
        system=system_prompt,
        messages=[{"role": "user", "content": user_prompt}],
    )
    return parse_json_response(response.content[0].text)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=str, default=str(DEFAULT_XLSX),
                        help=f"Path to PO workbook (default: {DEFAULT_XLSX})")
    parser.add_argument("--skip-parse", action="store_true",
                        help="Skip Excel parsing, only run LLM enrichment on existing pos collection")
    parser.add_argument("--skip-llm", action="store_true",
                        help="Skip LLM enrichment, only parse Excel into pos collection")
    parser.add_argument("--force", action="store_true",
                        help="Re-extract POs that already have extracted data")
    parser.add_argument("--dry-run", action="store_true",
                        help="Print the first PO prompt and exit (no API calls)")
    args = parser.parse_args()

    load_dotenv()
    mongo_uri = os.getenv("MONGO_URI")
    db_name = os.getenv("MONGO_DB", "trades_quotes")
    api_key = os.getenv("ANTHROPIC_API_KEY")

    if not mongo_uri:
        print("ERROR: MONGO_URI not set in .env"); sys.exit(1)

    print(f"Connecting to MongoDB ({db_name})...")
    db = MongoClient(mongo_uri)[db_name]

    # Customer lookup for linking
    customers = list(db.customers.find(
        {}, {"_id": 0, "customer_id": 1, "email": 1}
    ))
    customers_by_email = {
        c["email"].lower(): c for c in customers if c.get("email")
    }

    # ===== STAGE 1: Parse Excel =====
    if not args.skip_parse:
        xlsx_path = Path(args.xlsx)
        pos = parse_workbook(xlsx_path)
        print(f"  Parsed {len(pos)} POs from workbook\n")

        # Match each PO's customer_email to customer_id
        for po in pos:
            po["matched_customer_id"] = link_to_customer(po, customers_by_email)

        # Upsert into pos collection
        print("Saving to MongoDB pos collection...")
        for po in pos:
            db.pos.update_one(
                {"po_number": po["po_number"]},
                {"$set": {**po, "_parsed_at": time.time()}},
                upsert=True,
            )

        # Index for fast lookup
        db.pos.create_index("po_number", unique=True)
        db.pos.create_index("matched_customer_id")
        db.pos.create_index("po_status")

        print(f"  pos collection: {db.pos.count_documents({})} documents")
        matched = db.pos.count_documents({"matched_customer_id": {"$ne": None}})
        print(f"  POs linked to customers: {matched}/{db.pos.count_documents({})}")
    else:
        print("Skipping Excel parse (--skip-parse)\n")

    # ===== STAGE 2: LLM enrichment =====
    if args.skip_llm:
        print("\nSkipping LLM enrichment (--skip-llm)")
        return

    if not api_key and not args.dry_run:
        print("ERROR: ANTHROPIC_API_KEY not set in .env"); sys.exit(1)

    print("\nLoading POs for LLM enrichment...")
    query = {} if args.force else {"extracted": {"$exists": False}}
    pos_to_extract = list(db.pos.find(query))
    print(f"  POs to enrich: {len(pos_to_extract)}")

    if not pos_to_extract:
        print("Nothing to enrich. Use --force to re-extract.")
        return

    if args.dry_run:
        sample_prompt = build_po_prompt(pos_to_extract[0])
        print("\n" + "=" * 60)
        print("DRY RUN — first PO prompt")
        print("=" * 60)
        print("\n--- SYSTEM ---\n" + SYSTEM_PROMPT_POS)
        print("\n--- USER ---\n" + sample_prompt)
        return

    client = Anthropic(api_key=api_key)
    success = 0
    failures: list[tuple[str, str]] = []

    for po in tqdm(pos_to_extract, desc="Enriching"):
        prompt = build_po_prompt(po)
        try:
            extracted = extract_one(client, SYSTEM_PROMPT_POS, prompt)
            extracted["_extracted_at"] = time.time()
            extracted["_model"] = MODEL
            db.pos.update_one(
                {"po_number": po["po_number"]},
                {"$set": {"extracted": extracted}},
            )
            success += 1
        except Exception as e:
            failures.append((po["po_number"], str(e)[:200]))

    # Summary
    print(f"\n{'=' * 50}")
    print(f"Successful enrichments: {success}")
    print(f"Failures:               {len(failures)}")
    if failures:
        print("\nFirst 5 failures:")
        for pn, err in failures[:5]:
            print(f"  {pn}: {err}")

    # Sanity check
    print(f"\nVerifying in MongoDB...")
    total = db.pos.count_documents({})
    extracted_count = db.pos.count_documents({"extracted": {"$exists": True}})
    print(f"  POs with extracted field: {extracted_count}/{total}")

    # Distribution checks
    print(f"\n  Trade distribution (LLM-inferred):")
    pipeline = [
        {"$match": {"extracted": {"$exists": True}}},
        {"$group": {"_id": "$extracted.trade_inferred", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}},
    ]
    for r in db.pos.aggregate(pipeline):
        print(f"    {str(r['_id']):15} {r['count']:>5}")

    print(f"\n  Complexity distribution:")
    pipeline = [
        {"$match": {"extracted": {"$exists": True}}},
        {"$group": {"_id": "$extracted.complexity", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}},
    ]
    for r in db.pos.aggregate(pipeline):
        print(f"    {str(r['_id']):15} {r['count']:>5}")

    flagged = db.pos.count_documents(
        {"extracted.risk_flags": {"$exists": True, "$ne": []}}
    )
    print(f"\n  POs with risk flags: {flagged}/{total}")


if __name__ == "__main__":
    main()
